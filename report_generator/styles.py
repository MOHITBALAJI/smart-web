import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def apply_section_margins(section):
    """
    Enforces strictly equal margins on left and right as requested by user:
    Left: 1.0 inch
    Right: 1.0 inch
    Top: 1.0 inch
    Bottom: 1.0 inch
    """
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

def setup_document_styles(doc):
    for section in doc.sections:
        apply_section_margins(section)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def add_footer_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._p.append(fldSimple)

def add_title_paragraph(doc, text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)  # Pure black
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)  # Pure black
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)  # Pure black
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)  # Pure black
    return p

def add_body_p(doc, text, bold_prefix=None, indent_tab=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if indent_tab:
        p.paragraph_format.first_line_indent = Inches(0.5)  # 1 tab space indent
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text, bold_prefix=None, indent_level=1):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.25 * indent_level)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_clean_toc_table(doc, headers, rows):
    """
    Clean, borderless Table of Contents matching university photo.
    """
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [1.3, 4.2, 1.0]

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = headers[0]  # CHAPTER NO
    hdr_cells[1].text = headers[1]  # TITLE
    hdr_cells[2].text = headers[2]  # PAGE NO

    for i, c in enumerate(hdr_cells):
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 2 else WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        chap_no, title, page_no = row_data[0], row_data[1], row_data[2]
        
        is_chapter = chap_no.isdigit() or (chap_no == "" and title in ["ABSTRACT", "TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES"])
        
        row_cells[0].text = chap_no
        row_cells[1].text = title
        row_cells[2].text = str(page_no)

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(4)
        p0.paragraph_format.space_after = Pt(4)

        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(4)
        p1.paragraph_format.space_after = Pt(4)
        if not is_chapter and chap_no != "":
            p1.paragraph_format.left_indent = Inches(0.15)

        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(4)
        p2.paragraph_format.space_after = Pt(4)

        for p in [p0, p1, p2]:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
                if is_chapter:
                    r.font.bold = True

    # Remove borders
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders {} >'
                r'<w:top w:val="none"/>'
                r'<w:left w:val="none"/>'
                r'<w:bottom w:val="none"/>'
                r'<w:right w:val="none"/>'
                r'</w:tcBorders>'.format(nsdecls('w'))
            )
            tcPr.append(tcBorders)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(10)
    return table

def add_clean_list_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [1.3, 4.2, 1.0]

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = headers[0]
    hdr_cells[1].text = headers[1]
    hdr_cells[2].text = headers[2]

    for i, c in enumerate(hdr_cells):
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 2 else WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)

    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        row_cells[0].text = str(row_data[0])
        row_cells[1].text = str(row_data[1])
        row_cells[2].text = str(row_data[2])

        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)

        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)

        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(3)
        p2.paragraph_format.space_after = Pt(3)

        for p in [p0, p1, p2]:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(0, 0, 0)

    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders {} >'
                r'<w:top w:val="none"/>'
                r'<w:left w:val="none"/>'
                r'<w:bottom w:val="none"/>'
                r'<w:right w:val="none"/>'
                r'</w:tcBorders>'.format(nsdecls('w'))
            )
            tcPr.append(tcBorders)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(10)
    return table

def add_styled_table(doc, headers, rows, col_widths=None):
    """
    Renders academic table with pure black text in headers, light gray header fill,
    and clean single borders.
    """
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)  # Pure black header text
        # Clean subtle grayscale shading for header
        shading = parse_xml(r'<w:shd {} w:fill="EAEAEA"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        fill_color = "F9F9F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
            shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Apply Borders
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                r'<w:tcBorders {} >'
                r'<w:top w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
                r'<w:left w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
                r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
                r'<w:right w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
                r'</w:tcBorders>'.format(nsdecls('w'))
            )
            tcPr.append(tcBorders)

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    return table

def add_table_caption_below(doc, tbl_no, tbl_title):
    """
    Renders Table Caption AFTER / BELOW the table as requested by the user.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(f"Table {tbl_no}: {tbl_title}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)  # Pure black
    return p

def add_figure_placeholder(doc, fig_no, fig_title, height_inches=2.8):
    """
    Creates a clean, dedicated boxed frame placeholder so the user can easily paste
    their own image into Microsoft Word.
    """
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)

    # Clean border around placeholder
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders {} >'
        r'<w:top w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
        r'<w:left w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
        r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
        r'<w:right w:val="single" w:sz="6" w:space="0" w:color="666666"/>'
        r'</w:tcBorders>'.format(nsdecls('w'))
    )
    tcPr.append(tcBorders)
    shading = parse_xml(r'<w:shd {} w:fill="FAFAFA"/>'.format(nsdecls('w')))
    tcPr.append(shading)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(40)
    
    r1 = p.add_run(f"[ PLACEHOLDER: INSERT FIGURE {fig_no} HERE ]\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(80, 80, 80)
    
    r2 = p.add_run(f"({fig_title})")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(120, 120, 120)

    # Caption BELOW placeholder
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(14)
    run_cap = p_cap.add_run(f"Figure {fig_no}: {fig_title}")
    run_cap.font.name = 'Times New Roman'
    run_cap.font.size = Pt(11)
    run_cap.font.bold = True
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(0, 0, 0)
    return p_cap

def add_box_callout(doc, title, content_list):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    shading = parse_xml(r'<w:shd {} w:fill="F5F5F5"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    tcBorders = parse_xml(
        r'<w:tcBorders {} >'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="single" w:sz="24" w:space="0" w:color="000000"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>'.format(nsdecls('w'))
    )
    cell._tc.get_or_add_tcPr().append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    r_title = p.add_run(f"★ {title}\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(11.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    for item in content_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.25
        p_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_item = p_item.add_run(f"• {item}")
        r_item.font.name = 'Times New Roman'
        r_item.font.size = Pt(11)
        r_item.font.color.rgb = RGBColor(0, 0, 0)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)
